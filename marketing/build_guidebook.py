# Copyright 2026 Lakelet contributors
# SPDX-License-Identifier: Apache-2.0
"""Render marketing/guidebook.md to marketing/warehouse-for-one-guidebook.pptx.

    uv run --no-project --with python-pptx python marketing/build_guidebook.py

The markdown is the source of truth and is in git; the .pptx is generated and
gitignored, like deck/ and brand/. The format is deliberately small: slides are
separated by a line that is exactly ---; "# " is the title, "> " the lede,
"- " a bullet (with **bold** runs), "| " a table row (separator rows skipped),
"![alt](path)" an image on the right half, and any other line a note paragraph.
The palette is the site's dark theme (web/src/styles/palette.css).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "marketing" / "guidebook.md"
OUTPUT = ROOT / "marketing" / "warehouse-for-one-guidebook.pptx"
MARK = ROOT / "app" / "src-tauri" / "icons" / "icon.png"

BG, PANEL, PANEL2, LINE = "101F24", "15292E", "1B302E", "304247"
INK, INK_SOFT, MUTED, LIME, ON_LIME = "F0F2E8", "C9D5D0", "A5B8B9", "D6EE83", "172622"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
FONT = "Calibri"


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def parse(text: str) -> list[dict]:
    slides = []
    for chunk in re.split(r"^---$", text, flags=re.M):
        if not chunk.strip():
            continue
        slide = {"title": "", "lede": "", "items": []}
        for raw in chunk.splitlines():
            line = raw.rstrip()
            if not line.strip():
                continue
            if line.startswith("# "):
                slide["title"] = line[2:].strip()
            elif line.startswith("> "):
                slide["lede"] = line[2:].strip()
            elif line.startswith("- "):
                slide["items"].append(("bullet", line[2:].strip()))
            elif line.startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if all(re.fullmatch(r"-+", c) for c in cells):
                    continue
                if slide["items"] and slide["items"][-1][0] == "table":
                    slide["items"][-1][1].append(cells)
                else:
                    slide["items"].append(("table", [cells]))
            elif line.startswith("!["):
                m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
                slide["items"].append(("image", m.group(2)))
            else:
                slide["items"].append(("note", line.strip()))
        slides.append(slide)
    return slides


def add_runs(paragraph, text: str, size: int, colour: str, bold_colour: str = LIME) -> None:
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        run = paragraph.add_run()
        run.text = part
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = i % 2 == 1
        run.font.color.rgb = rgb(bold_colour if i % 2 == 1 else colour)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    return tf


def wrapped_lines(text: str, width, size: int) -> int:
    """How many lines PowerPoint will need: Calibri averages about 0.5 em per character."""
    chars_per_line = max(1, int(Emu(width).inches / (size * 0.5 / 72)))
    return max(1, -(-len(text.replace("**", "")) // chars_per_line))


def add_table(slide, rows: list[list[str]], x, y, w) -> int:
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    size = 11 if len(rows) > 7 else 12
    col_w = Emu(int(w / ncols))
    text_w = col_w - Inches(0.2)
    heights = []
    for r, row in enumerate(rows):
        lines = max(wrapped_lines(c, text_w, size + (1 if r == 0 else 0)) for c in row)
        heights.append(Inches(0.12) + Emu(int(Inches(1) * lines * (size + 1) * 1.25 / 72)))
    shape = slide.shapes.add_table(len(rows), ncols, x, y, w, Emu(sum(heights)))
    table = shape.table
    for c in range(ncols):
        table.columns[c].width = col_w
    for r, row in enumerate(rows):
        table.rows[r].height = heights[r]
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PANEL2 if r == 0 else PANEL)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            add_runs(p, text, size if r else size + 1, INK if r else LIME)
            if r == 0:
                for run in p.runs:
                    run.font.bold = True
    return Emu(sum(heights))


def render(slides: list[dict]) -> tuple[Presentation, list[tuple[str, float]]]:
    bottoms: list[tuple[str, float]] = []
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    for index, spec in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(BG)
        title_slide = index == 0
        has_image = any(kind == "image" for kind, _ in spec["items"])
        body_w = Inches(6.9) if has_image else W - 2 * MARGIN

        # A lime rule under the title on every slide but the first.
        if title_slide:
            tf = textbox(slide, MARGIN, Inches(2.1), W - 2 * MARGIN, Inches(1.2))
            add_runs(tf.paragraphs[0], spec["title"], 54, LIME)
            tf = textbox(slide, MARGIN, Inches(3.4), W - 2 * MARGIN, Inches(1.8))
            add_runs(tf.paragraphs[0], spec["lede"], 20, INK_SOFT)
            y = Inches(5.4)
            for kind, text in spec["items"]:
                if kind == "note":
                    tf = textbox(slide, MARGIN, y, W - 2 * MARGIN, Inches(0.5))
                    add_runs(tf.paragraphs[0], text, 14, MUTED)
                    y += Inches(0.45)
        else:
            tf = textbox(slide, MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(0.9))
            add_runs(tf.paragraphs[0], spec["title"], 32, LIME)
            rule = slide.shapes.add_shape(1, MARGIN, Inches(1.3), Inches(1.2), Inches(0.05))
            rule.fill.solid()
            rule.fill.fore_color.rgb = rgb(LIME)
            rule.line.fill.background()
            y = Inches(1.5)
            if spec["lede"]:
                tf = textbox(slide, MARGIN, y, body_w, Inches(0.6))
                add_runs(tf.paragraphs[0], spec["lede"], 18, INK_SOFT)
                y += Inches(0.6)
            for kind, payload in spec["items"]:
                if kind == "bullet":
                    tf = textbox(slide, MARGIN, y, body_w, Inches(0.9))
                    p = tf.paragraphs[0]
                    add_runs(p, "•  ", 15, LIME)
                    add_runs(p, payload, 15, INK)
                    lines = wrapped_lines(payload, body_w - Inches(0.4), 15)
                    y += Inches(0.3) * lines + Inches(0.1)
                elif kind == "table":
                    y += Inches(0.05)
                    y += add_table(slide, payload, MARGIN, y, body_w) + Inches(0.2)
                elif kind == "note":
                    tf = textbox(slide, MARGIN, y, body_w, Inches(0.6))
                    add_runs(tf.paragraphs[0], payload, 13, MUTED)
                    y += Inches(0.27) * wrapped_lines(payload, body_w, 13) + Inches(0.15)
                elif kind == "image":
                    path = ROOT / payload
                    if path.exists():
                        slide.shapes.add_picture(str(path), Inches(7.9), Inches(1.6), width=Inches(4.8))

        bottoms.append((spec["title"], round(Emu(y).inches, 2)))

        # Footer: the line, and the mark.
        tf = textbox(slide, MARGIN, H - Inches(0.55), Inches(9), Inches(0.4))
        add_runs(tf.paragraphs[0], f"A warehouse for one  ·  marketing/guidebook.md  ·  {index + 1} / {len(slides)}", 10, MUTED)
        if MARK.exists():
            slide.shapes.add_picture(str(MARK), W - MARGIN - Inches(0.35), H - Inches(0.55), height=Inches(0.35))
    return prs, bottoms


def main() -> int:
    slides = parse(SOURCE.read_text())
    prs, bottoms = render(slides)
    prs.save(OUTPUT)
    limit = Emu(H - Inches(0.7)).inches
    print(f"{OUTPUT.relative_to(ROOT)}: {len(slides)} slides; content must end above {limit:.2f} in")
    over = 0
    for i, (title, bottom) in enumerate(bottoms, 1):
        flag = "  OVER" if bottom > limit else ""
        over += bool(flag)
        print(f"  {i:2d}  {bottom:5.2f} in  {title[:58]}{flag}")
    return 1 if over else 0


if __name__ == "__main__":
    sys.exit(main())
